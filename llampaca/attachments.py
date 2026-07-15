"""
File attachments for the interactive chat session (the `/attach` command).

Phase 1 of the attachment feature: extract the text of a user-chosen
document (PDF, Word, or plain text) and inject it *directly* into the
conversation, so the model can answer questions about it. There is no
retrieval/RAG here yet — that is phase 2, planned for documents too large
to fit in the context window. Because of that, this module enforces a hard
size budget: an attachment may take at most ATTACH_MAX_CONTEXT_FRACTION of
the model's context window, and anything bigger is rejected with a clear
explanation instead of being silently truncated (a truncated contract or
paper is worse than an honest refusal — the model would answer confidently
from half a document).

Design notes
------------
- This is deliberately NOT a model tool. `/attach` is user-initiated: the
  user picks the file, the CLI extracts the text in Python, and the model
  receives it as part of the next user message. No tool registration, no
  extra prompt overhead for the (small) models — the tool-based variant
  (`read_pdf`) can be added later if model-initiated reading turns out to
  be needed.

- The extracted text is NOT appended to the history as its own message.
  Gemma's chat template hard-rejects two consecutive `user` messages
  ("Conversation roles must alternate user/assistant/..."), so a separate
  attachment message followed by the user's question would 400 on every
  prompt-mode model. Instead the CLI *stages* attachments and merges them
  into the next user message (see build_attachment_block / the /attach
  handling in cli.py) — one user turn containing document + question, which
  is valid for every template.

- Unlike the agent's file tools, `/attach` is NOT sandboxed to the
  workspace. The sandbox exists to contain the *model*; `/attach` arguments
  are typed by the user themself, who is already the authority over their
  own files (~/Documents/report.pdf is the typical case).

- pypdf / python-docx are imported lazily inside the extractors so that
  `llampaca --help` and sessions that never attach anything do not pay the
  import cost, and so that a broken/missing optional dependency only
  affects the /attach command instead of the whole CLI.
"""

from pathlib import Path

# Reuse the agent's chars-per-token heuristic so the budget math here and
# the trimming math in the agent loop are based on the same estimate.
from llampaca.agent.loop import CHARS_PER_TOKEN

# Maximum share of the model's context window a single /attach may occupy.
# Rationale for 0.35: the agent starts trimming history at the 0.80
# watermark, and the system prompt + a few conversation turns need room
# too. 35% leaves the document fully in context while the conversation
# about it still has the majority of the window to breathe. Documents that
# do not fit belong to phase 2 (RAG), not to a bigger fraction.
ATTACH_MAX_CONTEXT_FRACTION = 0.35

# Plain-text extensions read verbatim (no extraction library needed).
# Kept deliberately explicit: an unknown binary format must fail loudly
# rather than be decoded to mojibake and silently fed to the model.
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".log", ".json",
                   ".xml", ".yaml", ".yml", ".toml", ".ini", ".py", ".js",
                   ".ts", ".html", ".css", ".sh", ".rst", ".tex"}

# Formats that need a dedicated extractor.
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}


class AttachmentError(Exception):
    """
    A user-facing problem with an attachment (unsupported format, encrypted
    PDF, no extractable text, ...). The CLI shows the message verbatim, so
    every raise site phrases it as advice for the user, not as a traceback.
    """


def extract_text(path: Path) -> str:
    """
    Extract the plain text of a document, dispatching on the file extension.

    Args:
        path: File to extract. Must exist and be a regular file.

    Returns:
        The extracted text (never empty — an empty extraction raises).

    Raises:
        AttachmentError: missing file, unsupported format, or a format-
            specific failure (encrypted PDF, scanned PDF with no text
            layer, corrupt file, undecodable "text" file...).
    """
    if not path.exists():
        raise AttachmentError(f"File not found: {path}")
    if not path.is_file():
        raise AttachmentError(f"Not a regular file: {path}")

    suffix = path.suffix.lower()
    if suffix in PDF_EXTENSIONS:
        text = _extract_pdf(path)
    elif suffix in DOCX_EXTENSIONS:
        text = _extract_docx(path)
    elif suffix in TEXT_EXTENSIONS:
        text = _extract_plain_text(path)
    elif suffix == ".doc":
        # Old binary Word format: python-docx cannot read it, and saying
        # "unsupported" without the fix would leave the user stuck.
        raise AttachmentError(
            "Legacy .doc files are not supported. Re-save the document as "
            ".docx (File > Save As in Word) and attach that."
        )
    else:
        supported = ", ".join(
            sorted(PDF_EXTENSIONS | DOCX_EXTENSIONS | TEXT_EXTENSIONS)
        )
        raise AttachmentError(
            f"Unsupported file type '{suffix or '(no extension)'}'. "
            f"Supported: {supported}"
        )

    if not text.strip():
        # Reachable e.g. for a genuinely empty .txt/.docx; the PDF path has
        # its own, more specific empty-text error (scanned PDF hint).
        raise AttachmentError(
            f"'{path.name}' contains no extractable text."
        )
    return text


def _extract_plain_text(path: Path) -> str:
    """Read a plain-text file as UTF-8, tolerating stray invalid bytes."""
    try:
        # errors="replace" because real-world "text" files often carry a few
        # legacy-encoding bytes; losing those characters is acceptable for
        # chat context, aborting the whole attachment is not.
        return path.read_text(encoding="utf-8", errors="replace")
    except OSError as e:
        raise AttachmentError(f"Could not read '{path.name}': {e}") from e


def _extract_pdf(path: Path) -> str:
    """
    Extract the text layer of a PDF with pypdf.

    Pages are joined with visible `--- Page N ---` markers so the model can
    reference locations ("the table on page 3") and the user can follow.
    """
    # Lazy import (see module docstring). Wrapped so that a missing package
    # becomes a friendly in-session message with the fix, instead of a
    # ModuleNotFoundError that would crash the whole chat — the typical
    # case being deps installed in one Python environment while llampaca
    # runs from another (e.g. a conda env installed before this feature).
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as e:
        raise AttachmentError(
            "PDF support requires the 'pypdf' package, which is not "
            "installed in the Python environment llampaca is running from. "
            "Install it with: pip install pypdf"
        ) from e

    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            # An empty user password is common (owner-locked but readable
            # PDFs); try it before giving up.
            try:
                if not reader.decrypt(""):
                    raise AttachmentError(
                        f"'{path.name}' is password-protected. Remove the "
                        "password and attach it again."
                    )
            except AttachmentError:
                raise
            except Exception as e:
                raise AttachmentError(
                    f"'{path.name}' is encrypted and could not be opened: {e}"
                ) from e

        page_texts = []
        for number, page in enumerate(reader.pages, start=1):
            # extract_text() returns None/"" for pages without a text layer
            # (pure images); keep the marker anyway so page numbers stay
            # aligned with the original document.
            page_texts.append(
                f"--- Page {number} ---\n{(page.extract_text() or '').strip()}"
            )
    except AttachmentError:
        raise
    except (PdfReadError, OSError, ValueError) as e:
        raise AttachmentError(f"Could not parse '{path.name}' as PDF: {e}") from e

    text = "\n\n".join(page_texts)
    # A PDF whose every page extracted to "" is almost certainly a scan
    # (images of text, no text layer). Say so explicitly: "no text" without
    # the why would read like a bug to the user.
    if not "".join(t.split("---", 2)[-1] for t in page_texts).strip():
        raise AttachmentError(
            f"'{path.name}' has no text layer (scanned document?). "
            "OCR is not supported yet."
        )
    return text


def _extract_docx(path: Path) -> str:
    """
    Extract the text of a .docx with python-docx.

    Phase-1 scope: paragraph text plus a simple pipe-separated rendering of
    tables. Headers/footers, footnotes and text boxes are not extracted —
    acceptable for chat context, documented here so nobody wonders.
    """
    # Lazy import wrapped for the same reason as in _extract_pdf: a missing
    # package must degrade to an actionable message, never crash the session.
    try:
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError
    except ImportError as e:
        raise AttachmentError(
            "Word support requires the 'python-docx' package, which is not "
            "installed in the Python environment llampaca is running from. "
            "Install it with: pip install python-docx"
        ) from e

    try:
        document = Document(str(path))
    except (PackageNotFoundError, KeyError, OSError, ValueError) as e:
        raise AttachmentError(
            f"Could not parse '{path.name}' as a Word document: {e}"
        ) from e

    parts = [p.text for p in document.paragraphs]
    # Tables are separate from paragraphs in python-docx. Rendering them
    # after the paragraphs loses their exact position in the flow, but for
    # question-answering the content matters far more than the placement.
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def estimate_tokens(text: str) -> int:
    """Heuristic token estimate (chars/4), same basis as the agent loop."""
    return len(text) // CHARS_PER_TOKEN


def attachment_token_budget(context_size: int) -> int:
    """The maximum estimated tokens an attachment may occupy, given -c."""
    return int(context_size * ATTACH_MAX_CONTEXT_FRACTION)


def build_attachment_block(filename: str, text: str) -> str:
    """
    Frame an extracted document for injection into a user message.

    The begin/end markers matter: the model must be able to tell where the
    (untrusted, possibly instruction-looking) document stops and the user's
    actual request begins, and with multiple attachments where one document
    ends and the next starts.

    The "full content is included below" hint is there for the small local
    models: verified empirically that without it they reach for read_file/
    search_text to look for the attached file in the workspace (where it
    may not even exist) instead of just reading the text they were given.
    """
    return (
        f"[Attached file: {filename} — its full content is included below; "
        "read it from here, do NOT use tools to open it]\n"
        f"{text}\n"
        f"[End of attached file: {filename}]"
    )
